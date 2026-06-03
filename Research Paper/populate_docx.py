"""
populate_docx.py  —  Fill the MediCare AI Final Project Report template
                     with content from the 7 chapter Markdown files.

Usage:
    cd "Research Paper"
    python populate_docx.py

Output: MediCare_AI_Final_Report.docx
"""

import re
import copy
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import lxml.etree as etree

TEMPLATE_PATH = r"Template-07 - Final Project Report (1).docx"
OUTPUT_PATH   = r"MediCare_AI_Final_Report.docx"

# ── Chapter MD files ──────────────────────────────────────────────────────────
CHAPTER_FILES = {
    1: "Chapter_1_Introduction.md",
    2: "Chapter_2_SRS.md",
    3: "Chapter_3_SystemAnalysis.md",
    4: "Chapter_4_SystemDesign.md",
    5: "Chapter5.md",
    6: "Chapter6.md",
    7: "Chapter_7_Conclusion.md",
}

# ─────────────────────────────────────────────────────────────────────────────
# Helper: delete a paragraph element from the document body
# ─────────────────────────────────────────────────────────────────────────────
def _delete_paragraph(para):
    p = para._element
    p.getparent().remove(p)

# ─────────────────────────────────────────────────────────────────────────────
# Helper: insert a paragraph BEFORE a reference element
# ─────────────────────────────────────────────────────────────────────────────
def _insert_paragraph_before(doc, ref_para, text, style="Normal"):
    new_p = OxmlElement("w:p")
    ref_para._element.addprevious(new_p)
    new_para = doc.paragraphs[0]  # placeholder — we use low-level XML
    new_para._element = new_p
    return new_para

# ─────────────────────────────────────────────────────────────────────────────
# Apply standard body formatting to a paragraph
# ─────────────────────────────────────────────────────────────────────────────
def _fmt_body(para, size_pt=12, bold=False, italic=False, font_name="Times New Roman"):
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in para.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
    # line spacing 1.5 via pPr
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "360")        # 360 = 1.5 × 240 twips
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)

# ─────────────────────────────────────────────────────────────────────────────
# Parse a Markdown file and return a list of (type, content) tuples
#
# Types:  "h1" "h2" "h3" "h4" "code" "table" "blank" "body"
# ─────────────────────────────────────────────────────────────────────────────
def parse_md(path):
    with open(path, encoding="utf-8") as f:
        raw = f.read()

    # Strip formatting-hint lines left over from the MD template
    raw = re.sub(r"\*\*\[Font:.*?\]\*\*\n?", "", raw)

    lines = raw.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if line.startswith("```"):
            lang = line[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code_lines)))
            i += 1
            continue

        # Heading levels
        if line.startswith("#### "):
            blocks.append(("h4", line[5:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("---"):
            pass  # horizontal rule — skip
        elif line.startswith("|"):
            # Collect table rows
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            blocks.append(("table", "\n".join(table_lines)))
            continue
        elif line.strip() == "":
            blocks.append(("blank", ""))
        else:
            blocks.append(("body", line))

        i += 1

    return blocks


# ─────────────────────────────────────────────────────────────────────────────
# Parse a Markdown table → list of row lists
# ─────────────────────────────────────────────────────────────────────────────
def parse_table(raw):
    rows = []
    for line in raw.split("\n"):
        if re.match(r"^\|[-: |]+\|$", line.strip()):
            continue  # separator line
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Clean inline markdown bold/italic markers from a string
# ─────────────────────────────────────────────────────────────────────────────
def clean_inline(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*",   r"\1", text)
    text = re.sub(r"`(.+?)`",     r"\1", text)
    return text


# ─────────────────────────────────────────────────────────────────────────────
# Add a run with inline bold/italic detection
# ─────────────────────────────────────────────────────────────────────────────
def _add_inline_runs(para, text, default_bold=False, default_italic=False,
                     font_name="Times New Roman", size_pt=12):
    # Split on bold (**text**) and italic (*text*)
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        run = para.add_run()
        if tok.startswith("**") and tok.endswith("**"):
            run.text = tok[2:-2]
            run.bold = True
            run.italic = default_italic
        elif tok.startswith("*") and tok.endswith("*"):
            run.text = tok[1:-1]
            run.bold = default_bold
            run.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            run.text = tok[1:-1]
            run.font.name = "Courier New"
        else:
            run.text = tok
            run.bold = default_bold
            run.italic = default_italic
        run.font.name = font_name if not (tok.startswith("`") and tok.endswith("`")) else "Courier New"
        run.font.size = Pt(size_pt)


# ─────────────────────────────────────────────────────────────────────────────
# Append all blocks from a parsed MD file into a docx document
# (appends to end of document)
# ─────────────────────────────────────────────────────────────────────────────
def append_blocks(doc, blocks, skip_first_h1=True):
    first_h1_seen = False
    prev_blank = False

    for btype, bcontent in blocks:
        # Skip formatting hint lines
        if "[Font:" in bcontent:
            continue
        if "[Paragraph" in bcontent or "[Total Word" in bcontent:
            continue

        if btype == "h1":
            if skip_first_h1 and not first_h1_seen:
                first_h1_seen = True
                continue  # template already has the chapter heading
            first_h1_seen = True
            p = doc.add_paragraph(style="Heading 1")
            run = p.add_run(clean_inline(bcontent))
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.bold = True
            prev_blank = False

        elif btype == "h2":
            p = doc.add_paragraph(style="Heading 2")
            run = p.add_run(clean_inline(bcontent))
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.bold = True
            prev_blank = False

        elif btype == "h3":
            p = doc.add_paragraph(style="Heading 3")
            _add_inline_runs(p, bcontent, default_bold=True, default_italic=True,
                             font_name="Times New Roman", size_pt=12)
            prev_blank = False

        elif btype == "h4":
            p = doc.add_paragraph(style="Normal")
            _add_inline_runs(p, bcontent, default_bold=True,
                             font_name="Times New Roman", size_pt=12)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            prev_blank = False

        elif btype == "body":
            text = bcontent.strip()
            if not text:
                continue
            # Bullet-like lines
            if text.startswith("- ") or text.startswith("• "):
                p = doc.add_paragraph(style="List Bullet")
                _add_inline_runs(p, text[2:].strip(),
                                 font_name="Times New Roman", size_pt=12)
            elif re.match(r"^\d+\.\s", text):
                p = doc.add_paragraph(style="List Number")
                _add_inline_runs(p, re.sub(r"^\d+\.\s", "", text).strip(),
                                 font_name="Times New Roman", size_pt=12)
            else:
                if prev_blank or True:  # always new paragraph
                    p = doc.add_paragraph(style="Normal")
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    _add_inline_runs(p, text,
                                     font_name="Times New Roman", size_pt=12)
                    # 1.5 line spacing
                    pPr = p._p.get_or_add_pPr()
                    sp = OxmlElement("w:spacing")
                    sp.set(qn("w:line"), "360")
                    sp.set(qn("w:lineRule"), "auto")
                    pPr.append(sp)
            prev_blank = False

        elif btype == "code":
            # Render code as a preformatted block with Courier New 9pt
            for line in bcontent.split("\n"):
                p = doc.add_paragraph(style="No Spacing")
                run = p.add_run(line)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            prev_blank = False

        elif btype == "table":
            rows = parse_table(bcontent)
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=num_cols)
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    if ci >= num_cols:
                        break
                    cell = tbl.cell(ri, ci)
                    cell.text = clean_inline(cell_text)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)
                        if ri == 0:
                            for run in para.runs:
                                run.bold = True
            prev_blank = False

        elif btype == "blank":
            prev_blank = True


# ─────────────────────────────────────────────────────────────────────────────
# Find the index of the Nth chapter heading in doc.paragraphs
# Chapters start with a Heading 1 whose text is "Chapter N" (exact)
# ─────────────────────────────────────────────────────────────────────────────
def find_chapter_para_index(doc, chapter_num):
    target = f"Chapter {chapter_num}"
    for i, p in enumerate(doc.paragraphs):
        if p.style.name.startswith("Heading 1") and p.text.strip() == target:
            return i
    return -1


# ─────────────────────────────────────────────────────────────────────────────
# Delete all paragraphs between two indices (exclusive on both ends)
# ─────────────────────────────────────────────────────────────────────────────
def clear_chapter_content(doc, start_idx, end_idx):
    """Delete paragraphs from start_idx+1 up to (but not including) end_idx."""
    # Collect paragraph elements to delete
    to_delete = []
    paras = doc.paragraphs
    for i in range(start_idx + 1, end_idx):
        to_delete.append(paras[i]._element)

    for elem in to_delete:
        parent = elem.getparent()
        if parent is not None:
            parent.remove(elem)


# ─────────────────────────────────────────────────────────────────────────────
# Insert blocks into document body BEFORE the paragraph at end_idx
# (after clearing the chapter content)
# ─────────────────────────────────────────────────────────────────────────────
def insert_chapter_content(doc, end_para_elem, blocks, chapter_num):
    """
    Insert paragraphs generated from `blocks` immediately before `end_para_elem`
    in the document XML body.
    """
    body = doc.element.body

    skip_h1 = True  # skip the first H1 (chapter title — already in template)
    first_h1_seen = False
    prev_blank = False

    def _make_para_elem(style_name, text_parts, bold=False, italic=False,
                        font="Times New Roman", size=12, code=False,
                        spacing=True, align_just=True):
        """Build a low-level <w:p> element and insert it before end_para_elem."""
        p_elem = OxmlElement("w:p")

        # paragraph properties
        pPr = OxmlElement("w:pPr")
        p_elem.append(pPr)

        # style
        pStyle = OxmlElement("w:pStyle")
        # map style name to docx style id (must match the template's style IDs)
        style_id_map = {
            "Normal": "Normal",
            "Heading 1": "Heading1",
            "Heading 2": "Heading2",
            "Heading 3": "Heading3",
            "Heading3": "Heading3",
            "List Bullet": "ListBullet",
            "List Number": "ListNumber",
            "ListBullet": "ListBullet",
            "ListNumber": "ListNumber",
            "No Spacing": "NoSpacing",
            "NoSpacing": "NoSpacing",
        }
        pStyle.set(qn("w:val"), style_id_map.get(style_name, style_name))
        pPr.append(pStyle)

        # alignment
        if align_just:
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "both")
            pPr.append(jc)

        # 1.5 line spacing
        if spacing:
            sp_elem = OxmlElement("w:spacing")
            sp_elem.set(qn("w:line"), "360")
            sp_elem.set(qn("w:lineRule"), "auto")
            pPr.append(sp_elem)

        # run(s)
        for (txt, b, it, fn, sz) in text_parts:
            if not txt:
                continue
            r = OxmlElement("w:r")
            rPr = OxmlElement("w:rPr")
            rFonts = OxmlElement("w:rFonts")
            rFonts.set(qn("w:ascii"), fn)
            rFonts.set(qn("w:hAnsi"), fn)
            rPr.append(rFonts)
            sz_elem = OxmlElement("w:sz")
            sz_elem.set(qn("w:val"), str(sz * 2))
            rPr.append(sz_elem)
            if b:
                bold_e = OxmlElement("w:b")
                rPr.append(bold_e)
            if it:
                italic_e = OxmlElement("w:i")
                rPr.append(italic_e)
            r.append(rPr)
            t = OxmlElement("w:t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = txt
            r.append(t)
            p_elem.append(r)

        body.insert(list(body).index(end_para_elem), p_elem)
        return p_elem

    def _simple(text_parts, style="Normal", bold=False, italic=False,
                font="Times New Roman", size=12, spacing=True, align_just=True):
        parts = [(t, b, i, fn, sz) for (t, b, i, fn, sz) in text_parts]
        _make_para_elem(style, parts, bold, italic, font, size,
                        spacing=spacing, align_just=align_just)

    def txt_part(text, bold=False, italic=False,
                 font="Times New Roman", size=12):
        return (text, bold, italic, font, size)

    # ── parse inline markdown into run specs ──
    def parse_inline(text, def_bold=False, def_italic=False,
                     font="Times New Roman", size=12):
        tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
        parts = []
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                parts.append(txt_part(tok[2:-2], bold=True, italic=def_italic,
                                      font=font, size=size))
            elif tok.startswith("*") and tok.endswith("*"):
                parts.append(txt_part(tok[1:-1], bold=def_bold, italic=True,
                                      font=font, size=size))
            elif tok.startswith("`") and tok.endswith("`"):
                parts.append(txt_part(tok[1:-1], font="Courier New", size=size))
            else:
                parts.append(txt_part(tok, bold=def_bold, italic=def_italic,
                                      font=font, size=size))
        return parts

    # ── table helper ──
    def insert_table(rows_data):
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        tbl_elem = OxmlElement("w:tbl")

        # table properties
        tblPr = OxmlElement("w:tblPr")
        tblStyle = OxmlElement("w:tblStyle")
        tblStyle.set(qn("w:val"), "TableGrid")
        tblPr.append(tblStyle)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), "0")
        tblW.set(qn("w:type"), "auto")
        tblPr.append(tblW)
        tbl_elem.append(tblPr)

        # table grid
        tblGrid = OxmlElement("w:tblGrid")
        for _ in range(num_cols):
            gridCol = OxmlElement("w:gridCol")
            tblGrid.append(gridCol)
        tbl_elem.append(tblGrid)

        for ri, row_data in enumerate(rows_data):
            tr = OxmlElement("w:tr")
            for ci in range(num_cols):
                tc = OxmlElement("w:tc")
                tcPr = OxmlElement("w:tcPr")
                tc.append(tcPr)
                p_cell = OxmlElement("w:p")
                r_cell = OxmlElement("w:r")
                rPr_cell = OxmlElement("w:rPr")
                rFonts_c = OxmlElement("w:rFonts")
                rFonts_c.set(qn("w:ascii"), "Times New Roman")
                rFonts_c.set(qn("w:hAnsi"), "Times New Roman")
                rPr_cell.append(rFonts_c)
                sz_c = OxmlElement("w:sz")
                sz_c.set(qn("w:val"), "20")  # 10pt
                rPr_cell.append(sz_c)
                if ri == 0:
                    b_c = OxmlElement("w:b")
                    rPr_cell.append(b_c)
                r_cell.append(rPr_cell)
                t_cell = OxmlElement("w:t")
                t_cell.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                cell_text = row_data[ci] if ci < len(row_data) else ""
                t_cell.text = clean_inline(cell_text)
                r_cell.append(t_cell)
                p_cell.append(r_cell)
                tc.append(p_cell)
                tr.append(tc)
            tbl_elem.append(tr)

        body.insert(list(body).index(end_para_elem), tbl_elem)

    # ── main loop over blocks ──
    for btype, bcontent in blocks:
        # Skip MD formatting hint lines
        if "[Font:" in bcontent or "[Paragraph" in bcontent or "[Total Word" in bcontent:
            continue

        if btype == "h1":
            if skip_h1 and not first_h1_seen:
                first_h1_seen = True
                continue
            first_h1_seen = True
            parts = parse_inline(clean_inline(bcontent), def_bold=True,
                                 font="Times New Roman", size=14)
            _make_para_elem("Heading 1", parts, spacing=False, align_just=False)

        elif btype == "h2":
            parts = parse_inline(clean_inline(bcontent), def_bold=True,
                                 font="Times New Roman", size=13)
            _make_para_elem("Heading 2", parts, spacing=False, align_just=False)

        elif btype == "h3":
            parts = parse_inline(bcontent, def_bold=True, def_italic=True,
                                 font="Times New Roman", size=12)
            _make_para_elem("Heading 3", parts, spacing=True)

        elif btype == "h4":
            parts = parse_inline(bcontent, def_bold=True,
                                 font="Times New Roman", size=12)
            _make_para_elem("Normal", parts, spacing=True)

        elif btype == "body":
            text = bcontent.strip()
            if not text:
                continue
            if text.startswith("- ") or text.startswith("• "):
                parts = parse_inline(text[2:].strip(),
                                     font="Times New Roman", size=12)
                _make_para_elem("ListBullet", parts, spacing=False, align_just=False)
            elif re.match(r"^\d+\.\s", text):
                stripped = re.sub(r"^\d+\.\s", "", text).strip()
                parts = parse_inline(stripped, font="Times New Roman", size=12)
                _make_para_elem("ListNumber", parts, spacing=False, align_just=False)
            else:
                parts = parse_inline(text, font="Times New Roman", size=12)
                _make_para_elem("Normal", parts, spacing=True)

        elif btype == "code":
            for line in bcontent.split("\n"):
                _make_para_elem("NoSpacing",
                                [(line, False, False, "Courier New", 9)],
                                spacing=False, align_just=False)

        elif btype == "table":
            rows_data = parse_table(bcontent)
            insert_table(rows_data)

        elif btype == "blank":
            pass  # skip blank lines — spacing handled by paragraph spacing


# ─────────────────────────────────────────────────────────────────────────────
# EXTRA CONTENT: Features added after chapter MDs were written
# ─────────────────────────────────────────────────────────────────────────────
PRESCRIPTION_NOTIFICATION_ADDENDUM = """
## 4.2.4 Extended Domain: Prescription and Notification Entities

### Prescription

The Prescription entity was added to the domain after the initial design phase, once the clinical consultation workflow was fully implemented. Each prescription is uniquely linked to a completed appointment (one-to-one) and authored by the assigned doctor for the patient.

**Prescription Attributes:**
- **appointment_id** (ObjectID, Unique FK) — links prescription to the appointment in which it was issued; the unique constraint ensures one prescription per consultation
- **doctor_id** (ObjectID, FK) — references the issuing doctor
- **patient_id** (ObjectID, FK) — references the receiving patient
- **diagnosis** (String, Required) — free-text clinical diagnosis
- **medicines** (Array of sub-documents) — each entry contains: name, dosage, frequency, duration, instructions
- **vital_signs** (Sub-document) — blood pressure, temperature, pulse, weight recorded during consultation
- **advice** (String) — lifestyle or dietary advice
- **follow_up_date** (Date, Optional) — recommended date for follow-up visit
- **createdAt / updatedAt** (Date, Auto-timestamp)

### Notification

The Notification entity provides an in-app alert system for both doctors and patients. Notifications are created by the backend whenever a significant appointment event occurs (booking, cancellation, completion).

**Notification Attributes:**
- **recipient_id** (ObjectID, FK) — the user who receives the notification
- **recipient_role** (String Enum: doctor/patient) — identifies the user type
- **type** (String) — event category (e.g., appointment_booked, appointment_cancelled)
- **title** (String) — short notification title
- **message** (String) — detailed notification body
- **appointment_id** (ObjectID, Optional FK) — linked appointment if applicable
- **read** (Boolean, Default false) — marks whether the user has viewed the notification
- **createdAt** (Date, Auto-timestamp)
- **Compound Index:** (recipient_id, read, createdAt) — enables efficient unread-count queries

"""

PRESCRIPTION_PSEUDO_CODE = """
## 5.1.7 Prescription Creation Flow

Doctors issue prescriptions following a completed consultation. The backend validates ownership, creates the prescription record, and notifies the patient automatically.

```
PROCEDURE CreatePrescription(doctor_id, appointment_id, prescriptionData):

    appointment ← Appointment.findById(appointment_id)
                             .populate("patient_id doctor_id")

    IF appointment IS NULL THEN
        RETURN Error(404, "Appointment not found")
    END IF

    IF appointment.doctor_id._id ≠ doctor_id THEN
        RETURN Error(403, "Not authorised for this appointment")
    END IF

    IF appointment.status ≠ "completed" THEN
        RETURN Error(400, "Prescription can only be issued for completed appointments")
    END IF

    // Idempotency — one prescription per appointment
    existing ← Prescription.findOne({ appointment_id: appointment_id })
    IF existing IS NOT NULL THEN
        RETURN Error(409, "Prescription already issued for this appointment")
    END IF

    prescription ← Prescription.create({
        appointment_id : appointment_id,
        doctor_id      : doctor_id,
        patient_id     : appointment.patient_id._id,
        diagnosis      : prescriptionData.diagnosis,
        medicines      : prescriptionData.medicines,
        vital_signs    : prescriptionData.vital_signs,
        advice         : prescriptionData.advice,
        follow_up_date : prescriptionData.follow_up_date
    })

    // Notify patient
    Notification.create({
        recipient_id   : appointment.patient_id._id,
        recipient_role : "patient",
        type           : "prescription_issued",
        title          : "Prescription Ready",
        message        : "Dr. " + appointment.doctor_id.last_Name +
                         " has issued your prescription.",
        appointment_id : appointment_id
    })

    RETURN Success(prescription)

END PROCEDURE
```

## 5.1.8 Notification Polling Flow

Clients poll the notification endpoint on a short interval to display real-time alerts without requiring WebSocket infrastructure.

```
PROCEDURE PollNotifications(recipient_id, role):

    // Rate-limited endpoint — max 1 request per 10 seconds per client
    notifications ← Notification.find({
        recipient_id   : recipient_id,
        recipient_role : role
    }).sort({ createdAt: -1 }).limit(20)

    unread_count ← Notification.countDocuments({
        recipient_id : recipient_id,
        read         : FALSE
    })

    RETURN {
        notifications : notifications,
        unread_count  : unread_count
    }

END PROCEDURE

PROCEDURE MarkNotificationsRead(recipient_id):

    Notification.updateMany(
        { recipient_id: recipient_id, read: FALSE },
        { $set: { read: TRUE } }
    )

    RETURN Success({ marked_read: true })

END PROCEDURE

// Frontend polls every 30 seconds on dashboard pages
SCHEDULE PollNotifications EVERY 30 SECONDS (client-side setInterval)
```
"""

DOCTOR_APPROVAL_ADDENDUM_CH2 = """
## 2.X Doctor Credential Verification and Approval Requirements

### REQ-BE-DAUTH-001 — Doctor Registration with Credential Submission

**Priority:** HIGH | **Verification:** Test

Every new doctor registration request shall include: (a) a valid medical licence number (text field, required), and (b) a degree or licence document file (PDF, JPG, or PNG, maximum 10 MB, required). The system shall store the degree file in a dedicated server directory (`public/degrees/`) using Multer multipart upload middleware and record the filename in the Doctor document. Registration shall not issue a session token; instead, the doctor's `isApproved` field shall be explicitly set to `null` (pending), and the API response shall return `{ success: true, pending: true }` with a human-readable confirmation message.

### REQ-BE-DAUTH-002 — Tri-State Approval Status

**Priority:** HIGH | **Verification:** Test

The Doctor model shall maintain an `isApproved` field with three distinct semantic values:

- **null** — Registration is pending admin review. Doctor cannot log in.
- **true** — Registration approved by admin. Doctor may log in and access the dashboard.
- **false** — Registration rejected by admin. Doctor cannot log in and receives an explanatory message.

The field shall default to `true` in the database schema to ensure backward compatibility with pre-feature doctor accounts, while all new registrations explicitly set the value to `null` at creation time, overriding the schema default.

### REQ-BE-DAUTH-003 — Login Gate Enforcement

**Priority:** HIGH | **Verification:** Test

The doctor login endpoint (`POST /D_LogIn`) shall evaluate `isApproved` before issuing a JWT session token. If `isApproved === null`, the endpoint shall return HTTP 403 with `{ status: 0, approval_status: "pending", msg: "..." }`. If `isApproved === false`, the endpoint shall return HTTP 403 with `{ status: 0, approval_status: "rejected", msg: "..." }`. Only when `isApproved === true` shall the endpoint proceed to issue a JWT token and set the `doctorToken` httpOnly cookie.

### REQ-BE-DAUTH-004 — Admin Approval and Rejection Endpoints

**Priority:** HIGH | **Verification:** Test

The system shall expose three admin-protected endpoints:

- `GET /admin/pending-doctors` — Returns all Doctor documents where `isApproved === null`, sorted by registration date descending, excluding password hashes.
- `PATCH /admin/doctor/:id/approve` — Sets `isApproved = true` for the specified doctor.
- `PATCH /admin/doctor/:id/reject` — Sets `isApproved = false` for the specified doctor.

All three endpoints shall be protected by the Admin middleware and shall return `{ success: true, message: "...", data: doctor }` on success.

### REQ-BE-DAUTH-005 — Verified Doctor Visibility

**Priority:** HIGH | **Verification:** Test

The doctor browse/search endpoint (`GET /View_Doctor`) used by patients shall filter results to only include Doctor documents where `isApproved === true`. Pending and rejected doctors shall not appear in patient-facing search results, ensuring that patients only interact with administratively verified practitioners.

### REQ-FE-DAUTH-001 — Post-Registration Pending Screen

**Priority:** HIGH | **Verification:** Demonstration

Upon successful registration submission, the doctor signup page shall display a dedicated pending confirmation screen (replacing the registration form) that informs the doctor that: (a) the request has been forwarded for admin review, (b) licence and degree documents will be verified, and (c) the doctor will be notified via the login page upon approval. The screen shall not navigate to the dashboard or set any session role.

### REQ-FE-DAUTH-002 — Login Pending and Rejected States

**Priority:** HIGH | **Verification:** Demonstration

The doctor login page shall handle two additional HTTP 403 response states returned by the backend: (a) `approval_status: "pending"` — displays an amber-styled alert with a clock icon informing the doctor the request is under review, and (b) `approval_status: "rejected"` — displays a red-styled alert informing the doctor the registration was denied. These states shall be visually distinct from generic authentication errors.

### REQ-FE-DAUTH-003 — Admin Requests Tab

**Priority:** HIGH | **Verification:** Demonstration

The admin dashboard shall include a dedicated "Requests" navigation section displaying all pending doctor registrations. Each registration card shall show: doctor name, speciality, registration date, email, phone number, licence number, qualifications, and a link to open the uploaded degree document in a new browser tab. Each card shall provide Approve and Reject action buttons with per-button loading spinners and optimistic state removal on success. The sidebar navigation item shall display an amber badge with the count of pending registrations when any exist.

### REQ-FE-DAUTH-004 — Read-Only Credential Fields on Profile Edit

**Priority:** MEDIUM | **Verification:** Demonstration

The doctor profile edit page shall display the licence number and degree document as read-only fields in a visually distinct locked section. These fields shall not be editable by the doctor after registration, and the backend `Update_Profile` endpoint shall ignore any licence number or degree file values submitted through the profile update form.

"""

DOCTOR_APPROVAL_ADDENDUM_CH4 = """
## 4.2.5 Doctor Verification Entity

The Doctor Verification sub-system adds three new persistent fields to the Doctor document and introduces the concept of tri-state approval status to the domain model.

**Extended Doctor Attributes (Verification-Related):**
- **licenseNumber** (String, Required at registration) — The doctor's official medical council licence number, submitted at registration and locked thereafter. Displayed read-only on the profile edit page and visible to administrators in the pending requests panel.
- **degreeFile** (String, Nullable) — Filename of the uploaded degree or licence document (PDF, JPG, or PNG). The file is stored in `public/degrees/` on the backend server and served via a dedicated static route `/degrees/:filename`. Default is `null`; set at registration and locked thereafter.
- **isApproved** (Boolean | null) — Tri-state field encoding the administrative verification decision. `null` = registration pending review; `true` = approved and active; `false` = rejected. Schema default is `true` to preserve backward compatibility with pre-feature doctor accounts; new registrations explicitly override to `null` at the controller level.

**Approval State Machine:**

The Doctor entity transitions through the following approval states:

```
[New Registration] --> isApproved: null (Pending)
        |
        v
  Admin Reviews Credentials
        |
   +----+----+
   |         |
   v         v
Approve    Reject
   |         |
   v         v
isApproved  isApproved
= true      = false
(Active)   (Denied)
```

Once set to `true` or `false` by an admin, the state does not automatically revert. An administrator may re-approve a previously rejected doctor by calling the approve endpoint again.

## 4.2.6 Admin Approval Workflow

The administrative approval workflow coordinates credential review between doctor registration and first login. The workflow involves three backend API endpoints, two new frontend components (pending screen and admin requests tab), and the extended Doctor model described above.

**Registration Flow (Doctor Side):**
1. Doctor completes signup form including licenseNumber (text) and degreeFile (file upload).
2. Backend receives multipart/form-data, validates required fields, stores degree file via Multer, creates Doctor document with `isApproved: null`, and returns `{ success: true, pending: true }` without issuing a token.
3. Frontend replaces the registration form with a pending confirmation screen — no session is created, no navigation to the dashboard occurs.

**Login Gate (Doctor Side):**
1. Doctor submits credentials to `POST /D_LogIn`.
2. Backend loads Doctor document and checks `isApproved` before password verification.
3. If `null`: return 403 with `approval_status: "pending"` → frontend shows amber alert.
4. If `false`: return 403 with `approval_status: "rejected"` → frontend shows red alert.
5. If `true`: proceed with password comparison, issue JWT token, set httpOnly cookie.

**Admin Review Flow:**
1. Admin navigates to "Requests" tab in the admin dashboard.
2. Dashboard fetches `GET /admin/pending-doctors` and renders a card per pending registration.
3. Each card displays doctor information, licence number, qualifications, and a link to the uploaded degree document (opened via `GET /degrees/:filename` static route).
4. Admin clicks Approve or Reject; frontend calls `PATCH /admin/doctor/:id/approve` or `reject`.
5. On success, the card is removed from the pending list optimistically, and the main doctor list and stats are refreshed.

"""

DOCTOR_APPROVAL_ADDENDUM_CH5 = """
## 5.1.9 Doctor Registration and Approval Flow

New doctor registrations are held in a pending state and routed through an administrative review before the doctor gains platform access. The following pseudo-code documents both the registration ingestion and the admin approval/rejection operations.

```
PROCEDURE RegisterDoctor(formData, profilePictureFile, degreeFile):

    // Validate required credential fields
    IF formData.licenseNumber IS EMPTY THEN
        RETURN Error(400, "Licence number is required")
    END IF

    IF degreeFile IS NULL THEN
        RETURN Error(400, "Degree/licence document is required")
    END IF

    // Check for duplicate email
    existing ← Doctor.findOne({ email: formData.email })
    IF existing IS NOT NULL THEN
        RETURN Error(400, "Email already registered")
    END IF

    // Store uploaded files via Multer
    profilePicFilename ← profilePictureFile?.filename OR "default-doctor.png"
    degreeFilename     ← degreeFile.filename

    // Parse degrees array from JSON string
    degreesArray ← JSON.parse(formData.degrees)

    // Create Doctor document — isApproved explicitly null (pending)
    doctor ← Doctor.create({
        first_Name    : formData.first_Name,
        last_Name     : formData.last_Name,
        email         : formData.email,
        password      : formData.password,     // hashed by pre-save hook
        ph            : formData.ph,
        speciality    : formData.speciality,
        degrees       : degreesArray,
        licenseNumber : formData.licenseNumber,
        profile_Picture : profilePicFilename,
        degreeFile    : degreeFilename,
        isApproved    : null                   // override schema default of true
    })

    // Do NOT issue token — return pending confirmation only
    RETURN Success(200, {
        success : true,
        pending : true,
        message : "Registration submitted. Await admin approval."
    })

END PROCEDURE


PROCEDURE DoctorLogin(email, password):

    user ← Doctor.findOne({ email: email })
    IF user IS NULL THEN
        RETURN Error(401, "Invalid credentials")
    END IF

    // Approval gate — checked BEFORE password verification
    IF user.isApproved === null THEN
        RETURN Error(403, {
            status          : 0,
            approval_status : "pending",
            msg             : "Your registration is under review. Kindly wait."
        })
    END IF

    IF user.isApproved === false THEN
        RETURN Error(403, {
            status          : 0,
            approval_status : "rejected",
            msg             : "Your registration request was denied."
        })
    END IF

    // Proceed with authentication only for approved doctors
    passwordMatch ← bcrypt.compare(password, user.password)
    IF NOT passwordMatch THEN
        RETURN Error(401, "Invalid credentials")
    END IF

    token ← JWT.sign({ id: user._id }, SECRET_KEY, { expiresIn: EXPIRE_IN })
    SET httpOnly cookie "doctorToken" = token

    RETURN Success(200, { status: 1, msg: "Login Successful", doctor: user })

END PROCEDURE


PROCEDURE AdminApproveDoctor(adminId, doctorId):

    doctor ← Doctor.findByIdAndUpdate(
        doctorId,
        { isApproved: true },
        { new: true }
    ).select("-password")

    IF doctor IS NULL THEN
        RETURN Error(404, "Doctor not found")
    END IF

    RETURN Success(200, {
        success : true,
        message : "Dr. " + doctor.first_Name + " " + doctor.last_Name + " approved.",
        data    : doctor
    })

END PROCEDURE


PROCEDURE AdminRejectDoctor(adminId, doctorId):

    doctor ← Doctor.findByIdAndUpdate(
        doctorId,
        { isApproved: false },
        { new: true }
    ).select("-password")

    IF doctor IS NULL THEN
        RETURN Error(404, "Doctor not found")
    END IF

    RETURN Success(200, {
        success : true,
        message : "Dr. " + doctor.first_Name + " " + doctor.last_Name + " rejected.",
        data    : doctor
    })

END PROCEDURE
```

"""

CHAPTER7_PRESCRIPTION_UPDATE = """
## 7.1.2 Scope and Deliverables (Updated)

**In-Scope Deliverables (Completed in v1.0):**

- **Frontend Application:** React 19 + Vite web application with responsive design supporting patient and doctor portals, admin portal (port 5173)
- **Backend API:** Express 5 Node.js REST API with comprehensive route handlers, middleware, and business logic (port 4000)
- **AI Service:** FastAPI Python service implementing RAG pipeline, audio transcription, and consultation summarization (port 8000)
- **Database:** MongoDB with 9 collections (Patient, Doctor, Schedule, Appointment, Chat, MeetingNote, Review, Prescription, Notification)
- **Video Infrastructure:** VideoSDK.live integration for real-time audio/video consultations
- **Prescription Management:** Doctor-issued digital prescriptions with medicines, dosage, vital signs, and patient notification
- **Notification System:** In-app alert system for both doctors and patients with unread count badge
- **Admin Portal:** Administrative dashboard for managing doctors, patients, appointments, and platform analytics, including a dedicated doctor registration review panel
- **Doctor Credential Verification System:** Admin-gated approval workflow requiring licence number and degree document submission at registration; tri-state approval status (pending/approved/rejected) enforced at login; only verified doctors appear in patient search results
- **Doctor Patient History:** Doctors can view a patient's full consultation history prior to prescription creation
- **Dashboard Analytics:** Real-time stats for doctors (total appointments, upcoming, completed, available slots) and admin
- **Documentation:** Complete specification across 7 chapters
- **Deployment:** Docker Compose configuration for complete stack deployment
- **Testing:** 144 tests across 8 test suites, all passing at 100%

"""


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────
def main():
    print(f"Opening template: {TEMPLATE_PATH}")
    doc = Document(TEMPLATE_PATH)

    print(f"Total paragraphs in template: {len(doc.paragraphs)}")

    # ── Locate each chapter boundary using ELEMENT REFERENCES (not indices) ──
    # This is critical: indices shift as we insert/delete content,
    # but XML element references remain stable.
    chapter_start_elems  = {}   # ch → XML element of "Chapter N" heading
    chapter_end_elems    = {}   # ch → XML element of first para of NEXT chapter (insert before)

    for ch in range(1, 8):
        idx = find_chapter_para_index(doc, ch)
        if idx >= 0:
            chapter_start_elems[ch] = doc.paragraphs[idx]._element
            print(f"  Chapter {ch} heading found at paragraph index {idx}: "
                  f"'{doc.paragraphs[idx].text[:60]}'")
        else:
            print(f"  WARNING: Chapter {ch} heading NOT found")

    # The "end" of chapter N is the start of chapter N+1 (or end marker)
    # Capture all these element references NOW before any modification
    sorted_chapters = sorted(chapter_start_elems.keys())
    for i, ch in enumerate(sorted_chapters):
        if i + 1 < len(sorted_chapters):
            next_ch = sorted_chapters[i + 1]
            chapter_end_elems[ch] = chapter_start_elems[next_ch]
        else:
            # Last chapter — find the Appendix / end matter element
            end_elem = None
            for p in doc.paragraphs:
                if p.style.name.startswith("Heading 1") and (
                    "Appendix" in p.text or "References" in p.text
                    or "Bibliography" in p.text
                ):
                    end_elem = p._element
                    break
            if end_elem is None:
                end_elem = doc.element.body[-1]  # sectPr
            chapter_end_elems[ch] = end_elem

    # ── Process each chapter ──────────────────────────────────────────────────
    for ch in sorted_chapters:
        md_file = CHAPTER_FILES[ch]
        start_elem = chapter_start_elems[ch]
        end_elem   = chapter_end_elems[ch]
        print(f"\nProcessing Chapter {ch} ({md_file}) ...")

        # Parse MD
        blocks = parse_md(md_file)
        print(f"  Parsed {len(blocks)} blocks from {md_file}")

        # For Chapter 2, append Doctor Approval requirements
        if ch == 2:
            extra_blocks = parse_md_text(DOCTOR_APPROVAL_ADDENDUM_CH2)
            blocks = blocks + extra_blocks

        # For Chapter 4, append Prescription/Notification content + Doctor Approval entity
        if ch == 4:
            extra_blocks = parse_md_text(PRESCRIPTION_NOTIFICATION_ADDENDUM)
            extra_blocks2 = parse_md_text(DOCTOR_APPROVAL_ADDENDUM_CH4)
            blocks = blocks + extra_blocks + extra_blocks2

        # For Chapter 5, append pseudo-codes including Doctor Approval flow
        if ch == 5:
            extra_blocks = parse_md_text(PRESCRIPTION_PSEUDO_CODE)
            extra_blocks2 = parse_md_text(DOCTOR_APPROVAL_ADDENDUM_CH5)
            blocks = blocks + extra_blocks + extra_blocks2

        # For Chapter 7, update scope/deliverables
        if ch == 7:
            extra_blocks = parse_md_text(CHAPTER7_PRESCRIPTION_UPDATE)
            blocks = extra_blocks + blocks

        # Delete all XML children of body that lie between start_elem and end_elem
        # (exclusive: keep start_elem, keep end_elem)
        body = doc.element.body
        body_children = list(body)
        start_i = body_children.index(start_elem)
        end_i   = body_children.index(end_elem)

        to_delete = body_children[start_i + 1 : end_i]
        print(f"  Clearing {len(to_delete)} existing elements between chapter markers ...")
        for elem in to_delete:
            body.remove(elem)

        # Insert new content before end_elem
        print(f"  Inserting new content ...")
        insert_chapter_content(doc, end_elem, blocks, ch)

        print(f"  Chapter {ch} done.")

    # ── Save ──────────────────────────────────────────────────────────────────
    print(f"\nSaving to: {OUTPUT_PATH}")
    doc.save(OUTPUT_PATH)
    print("Done! Document saved successfully.")


def parse_md_text(text):
    """Parse a raw markdown string (not a file) into blocks."""
    import tempfile, os
    with tempfile.NamedTemporaryFile(mode="w", suffix=".md",
                                     encoding="utf-8", delete=False) as tf:
        tf.write(text)
        tmp_path = tf.name
    blocks = parse_md(tmp_path)
    os.unlink(tmp_path)
    return blocks


if __name__ == "__main__":
    main()
