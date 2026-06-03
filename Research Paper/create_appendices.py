"""
Generates MediCare_AI_Appendices.docx — Appendices A, B, C, D
Formatting spec:
  Appendix title  : 20 pt, Calibri, Bold, Left
  Level-1 heading : 16 pt, Calibri, Bold, Left
  Level-2 heading : 14 pt, Calibri, Bold, Left
  Level-3 heading : 12 pt, Calibri, Bold, Left
  Body paragraph  : 12 pt, Calibri, 1.5 line spacing, Justified
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)


# ── Helpers ───────────────────────────────────────────────────────────────────

def set_line_spacing_15(para):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:line'),    '360')   # 240 = single, 360 = 1.5
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)


def add_page_break(doc):
    para = doc.add_paragraph()
    run  = para.add_run()
    run.add_break(docx.oxml.ns.qn and __import__('docx').enum.text.WD_BREAK if False else None)
    # simpler: use a paragraph with a page-break element
    from docx.oxml import OxmlElement
    br = OxmlElement('w:br')
    br.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'page')
    para._p.append(br)


def append_title(doc, text):
    """Appendix title — 20 pt, Calibri, Bold, Left"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(para)
    run = para.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
    return para


def append_h1(doc, text):
    """Level-1 heading — 16 pt, Calibri, Bold, Left"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(para)
    run = para.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
    return para


def append_h2(doc, text):
    """Level-2 heading — 14 pt, Calibri, Bold, Left"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(para)
    run = para.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
    return para


def append_h3(doc, text):
    """Level-3 heading — 12 pt, Calibri, Bold, Left"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_line_spacing_15(para)
    run = para.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
    return para


def append_body(doc, text):
    """Body paragraph — 12 pt, Calibri, 1.5 spacing, Justified"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_line_spacing_15(para)
    run = para.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x2E)
    return para


def append_spacer(doc):
    p = doc.add_paragraph()
    set_line_spacing_15(p)
    return p


import docx  # ensure import available for add_page_break internal


def page_break(doc):
    para = doc.add_paragraph()
    from docx.oxml import OxmlElement
    br = OxmlElement('w:br')
    br.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'page')
    para._p.append(br)


# ══════════════════════════════════════════════════════════════════════════════
# APPENDIX A — USER MANUAL
# ══════════════════════════════════════════════════════════════════════════════

append_title(doc, 'Appendix A: User Manual')
append_spacer(doc)
append_body(doc,
    'This appendix serves as a comprehensive operational guide for all end-users of the MediCare AI '
    'telemedicine platform, covering two primary user roles: patients and doctors. It provides '
    'step-by-step instructions for account registration, profile management, appointment booking, '
    'video consultation participation, AI-powered symptom assessment, and digital prescription access. '
    'Screenshots and navigational cues are described textually to guide users through each interface '
    'without prior technical knowledge. The manual is structured to follow the natural workflow a user '
    'would experience from first accessing the platform through completing a full teleconsultation cycle, '
    'ensuring a smooth and self-sufficient onboarding experience for all stakeholders.')
append_spacer(doc)

# ── A.1 Patient Guide ─────────────────────────────────────────────────────────
append_h1(doc, 'A.1  Patient User Guide')
append_body(doc,
    'This section documents every step a patient must follow to make full use of the MediCare AI '
    'platform, from initial registration through completing a teleconsultation and accessing their '
    'digital prescription.')
append_spacer(doc)

append_h2(doc, 'A.1.1  Registration and Login')
append_body(doc,
    'To register as a patient, navigate to the MediCare AI landing page and click the "Sign Up" '
    'button under the Patient section. The registration form collects the following mandatory fields: '
    'first name, last name, email address, password, date of birth, gender, and contact number. All '
    'fields are validated in real time; the email must be unique and the password must be at least '
    'eight characters. Upon successful submission, the system creates a patient account and redirects '
    'the user to the Patient Dashboard. Subsequent logins require only the registered email and '
    'password. Session authentication is maintained via a secure HTTP-only cookie that expires after '
    'a configurable period, eliminating the need to log in on every visit within the same session.')
append_spacer(doc)

append_h3(doc, 'A.1.1.1  Recovering a Forgotten Password')
append_body(doc,
    'If a patient forgets their password, they should contact the platform administrator through the '
    'support channel listed on the landing page. Direct password-reset functionality via email link '
    'is planned for a future release. In the interim, the administrator can reset credentials through '
    'the Admin Panel.')
append_spacer(doc)

append_h2(doc, 'A.1.2  Browsing Doctors and Booking an Appointment')
append_body(doc,
    'After logging in, patients land on the Patient Dashboard, which displays active appointments, '
    'quick statistics, and navigation links. To book an appointment, click "Find a Doctor" in the '
    'navigation bar or on the dashboard shortcut. The Doctors Listing page displays all registered '
    'and active doctors, each showing their name, speciality, experience, and average rating. Use '
    'the search bar to filter by name or speciality. Click a doctor card to view their full profile, '
    'qualifications, and available time slots. Select an available date from the calendar, choose a '
    'time slot, review the consultation fee, and click "Confirm Booking." The system automatically '
    'creates a VideoSDK meeting room and associates it with the appointment. A confirmation is '
    'displayed on screen and the appointment appears immediately under "My Appointments."')
append_spacer(doc)

append_h3(doc, 'A.1.2.1  Cancelling or Rescheduling an Appointment')
append_body(doc,
    'A patient may cancel a booked appointment at any time before it begins by navigating to '
    '"My Appointments" and clicking the "Cancel" button on the relevant appointment card. The system '
    'will prompt for confirmation before proceeding. If a doctor initiates a reschedule, the patient '
    'receives a free-rebook token displayed as a "Rebook Free" button on the cancelled appointment '
    'card. Clicking it opens the slot selection panel for the same doctor; no payment is required '
    'when redeeming this token.')
append_spacer(doc)

append_h2(doc, 'A.1.3  Joining a Video Consultation')
append_body(doc,
    'On the day of the appointment, navigate to "My Appointments." The appointment card shows a '
    'live countdown. The "Join Call" button becomes active five minutes before the scheduled start '
    'time and deactivates when the slot\'s end time is reached. Clicking "Join Call" redirects the '
    'patient to the secure video room powered by VideoSDK. The room supports HD video, two-way '
    'audio, screen sharing, and in-call chat. No plugin or software installation is required; the '
    'call runs entirely in a modern web browser. After the consultation ends, the backend '
    'automatically updates the appointment status to "Completed" and initiates the AI-powered '
    'meeting-notes pipeline.')
append_spacer(doc)

append_h3(doc, 'A.1.3.1  Downloading Meeting Notes')
append_body(doc,
    'Once a consultation is marked completed, the platform transcribes the audio recording using '
    'the faster-whisper model and generates a structured summary using the Groq LLM. When '
    'processing is finished — typically within two to five minutes — a "Download Meeting Notes (PDF)" '
    'button appears on the appointment card. Click it to download a patient-friendly PDF summary '
    'containing diagnosis points, key discussion items, and recommended follow-up actions.')
append_spacer(doc)

append_h2(doc, 'A.1.4  AI-Powered Symptom Assessment')
append_body(doc,
    'The AI Chat feature is accessible from the Patient Dashboard via the "AI Symptom Advisor" link '
    'in the navigation bar. The interface presents a conversational chat window. Type a symptom or '
    'health question in natural language and press "Send." The AI engine queries a curated knowledge '
    'base of five medical reference documents using a Retrieval-Augmented Generation (RAG) pipeline. '
    'Each response includes the AI-generated answer and cited source references, enabling patients to '
    'verify information. The chat history is preserved within the session. The AI advisor is strictly '
    'informational and does not replace professional medical consultation; a disclaimer is displayed '
    'at the top of the chat window at all times.')
append_spacer(doc)

append_h2(doc, 'A.1.5  Viewing Digital Prescriptions')
append_body(doc,
    'After a completed appointment, the attending doctor may issue a digital prescription through '
    'the platform. Patients receive an in-platform notification when a prescription is ready. Navigate '
    'to "My Appointments," locate the completed appointment card, and click "View Prescription." The '
    'prescription page displays the full diagnosis, a structured medicines table (name, dosage, '
    'frequency, duration, and special instructions), recorded vital signs, general advice, and the '
    'recommended follow-up date. The view is formatted for easy reading and can be printed directly '
    'from the browser using the standard print function (Ctrl+P).')
append_spacer(doc)

append_h2(doc, 'A.1.6  Submitting a Review')
append_body(doc,
    'Patients may rate and review their doctor after a completed appointment. On the completed '
    'appointment card, click "Leave a Review." A modal dialog opens with a five-star rating selector '
    'and an optional text field accepting up to 1,000 characters. Select the rating, optionally add '
    'comments, and click "Submit Review." Once submitted, the button is replaced by a "Reviewed" '
    'badge with a star icon. Reviews are publicly visible on the doctor\'s profile page and contribute '
    'to the doctor\'s aggregate rating displayed in the Doctors Listing.')
append_spacer(doc)

# ── A.2 Doctor Guide ──────────────────────────────────────────────────────────
append_h1(doc, 'A.2  Doctor User Guide')
append_body(doc,
    'This section provides step-by-step instructions for doctors using the MediCare AI platform '
    'to manage their availability, conduct video consultations, issue digital prescriptions, and '
    'review patient medical history.')
append_spacer(doc)

append_h2(doc, 'A.2.1  Doctor Registration and Profile Setup')
append_body(doc,
    'Navigate to the landing page and click "Doctor Sign Up." The registration form requires: '
    'first name, last name, email, password, medical speciality, years of experience, consultation '
    'fee, and a professional profile photograph. Optionally, a brief biography describing areas of '
    'expertise may be included. After submission the doctor is redirected to the Doctor Dashboard. '
    'Profile information can be updated at any time via "Edit Profile" in the navigation menu, '
    'including the profile photograph and consultation fee.')
append_spacer(doc)

append_h2(doc, 'A.2.2  Creating Availability Schedules')
append_body(doc,
    'Doctors control their bookable time slots through the "My Schedule" section. Click "Create '
    'Schedule" and select a date from the date picker. Add one or more time slots by specifying a '
    'start time and end time for each slot; overlapping slots are rejected by the system. Click '
    '"Save" to publish the schedule. Published slots immediately become visible to patients on the '
    'doctor\'s profile booking page. Slots can be deleted individually from the "My Schedule" view '
    'as long as no appointment has been booked against them.')
append_spacer(doc)

append_h3(doc, 'A.2.2.1  Managing Existing Schedules')
append_body(doc,
    'The "My Schedule" page lists all scheduled slots grouped by date. Each slot shows its date, '
    'start and end time, and current status (Available, Booked, or Completed). Slots with a '
    '"Booked" or "Completed" status cannot be deleted. Use the date filter to navigate between '
    'dates. The total slot count is displayed in the page header for quick reference.')
append_spacer(doc)

append_h2(doc, 'A.2.3  Managing Appointments')
append_body(doc,
    'The "My Appointments" page displays all patient appointments grouped by status. Doctors can '
    'filter appointments using the status bar at the top of the page: All, Upcoming, Ongoing, '
    'Completed, and Cancelled. Each appointment card shows the patient\'s name, scheduled date and '
    'time, and current status. For booked appointments, the "Reschedule" button allows the doctor '
    'to cancel the booking and issue the patient a free-rebook token. The "Start Consult" button '
    'becomes active five minutes before the slot starts and opens the VideoSDK video room.')
append_spacer(doc)

append_h2(doc, 'A.2.4  Writing Digital Prescriptions')
append_body(doc,
    'After a consultation is completed, click the "Prescription" button on the appointment card. '
    'The prescription modal opens with fields for: vital signs (blood pressure, temperature, pulse, '
    'and weight), diagnosis text, a dynamic medicines table where each medicine entry includes name, '
    'dosage, frequency, duration, and special instructions, a general advice text area, and a '
    'follow-up date picker. Additional medicine rows can be added using the "+ Add Medicine" button '
    'or removed using the row\'s close icon. Click "Save Prescription" to submit. The system '
    'enforces one prescription per appointment (idempotency); re-opening the modal loads the '
    'existing prescription for editing. The patient receives an in-platform notification when the '
    'prescription is saved.')
append_spacer(doc)

append_h2(doc, 'A.2.5  Viewing Patient History')
append_body(doc,
    'Click the "History" button on any appointment card to open the Patient History modal. The modal '
    'contains two tabs: Appointments and Prescriptions. The Appointments tab shows a detailed card '
    'for each past appointment including the attending doctor, date, time, status, and a reference '
    'ID. The Prescriptions tab shows all previously issued prescriptions for the patient, each '
    'displaying the diagnosis, vital signs recorded, full medicines list with all details, advice, '
    'and follow-up date. This enables doctors to make informed clinical decisions before issuing '
    'new prescriptions.')
append_spacer(doc)

page_break(doc)

# ══════════════════════════════════════════════════════════════════════════════
# APPENDIX B — ADMINISTRATOR MANUAL
# ══════════════════════════════════════════════════════════════════════════════

append_title(doc, 'Appendix B: Administrator Manual')
append_spacer(doc)
append_body(doc,
    'This appendix provides a complete operational reference for system administrators responsible '
    'for managing the MediCare AI platform. It covers access procedures for the Admin Dashboard, '
    'monitoring platform-wide statistics, and performing data management operations across all '
    'system entities including doctors, patients, appointments, schedules, meeting notes, and '
    'reviews. Administrators hold the highest privilege level within the application and are '
    'responsible for maintaining data integrity, moderating platform content, and ensuring '
    'operational continuity. This manual assumes the administrator has network access to the '
    'deployed application and holds valid administrator credentials issued during system setup.')
append_spacer(doc)

# B.1
append_h1(doc, 'B.1  Accessing the Admin Dashboard')
append_body(doc,
    'The Admin Dashboard is accessible at the dedicated route /admin/login, which is not linked '
    'from the main patient or doctor interfaces and must be accessed by direct URL navigation. '
    'Enter the administrator email and password in the login form and click "Sign In." Successful '
    'authentication sets a secure HTTP-only admin session cookie valid for the duration of the '
    'browser session. The dashboard is then accessible at /admin, presenting summary statistics '
    'and navigation links to each management section.')
append_spacer(doc)

append_h2(doc, 'B.1.1  Dashboard Statistics Overview')
append_body(doc,
    'The top section of the Admin Dashboard displays real-time aggregate statistics retrieved from '
    'the backend GET /admin/stats endpoint. The statistics panel shows: total registered doctors, '
    'total registered patients, total appointments (all statuses), total completed appointments, '
    'total cancelled appointments, total meeting notes generated, and total reviews submitted. '
    'These figures update each time the page is loaded or refreshed and provide a high-level view '
    'of platform health and usage.')
append_spacer(doc)

append_h3(doc, 'B.1.1.1  Session Management')
append_body(doc,
    'The admin session is maintained via an HTTP-only cookie. There is no automatic idle timeout '
    'within the current release; administrators should manually log out using the "Logout" button '
    'accessible in the admin navigation bar after completing their session. The backend POST '
    '/admin/logout endpoint clears the session cookie server-side, ensuring the token cannot be '
    'reused after logout.')
append_spacer(doc)

# B.2
append_h1(doc, 'B.2  Managing Doctors')
append_body(doc,
    'Navigate to the Doctors section from the Admin Dashboard sidebar. The page fetches all '
    'registered doctor records via GET /admin/doctors and displays them in a searchable, paginated '
    'table showing the doctor\'s name, email, speciality, experience, consultation fee, and '
    'registration date. Administrators can search by name or speciality using the search bar.')
append_spacer(doc)

append_h2(doc, 'B.2.1  Removing a Doctor Account')
append_body(doc,
    'To delete a doctor record, locate the doctor row in the table and click the "Delete" button. '
    'A confirmation dialog appears requiring explicit confirmation before the action proceeds. '
    'The backend executes DELETE /admin/doctor/:id, which removes the doctor document from the '
    'MongoDB Doctor collection. Note that appointments, schedules, and reviews associated with the '
    'deleted doctor remain in the database for audit purposes but will show the doctor as unavailable '
    'in patient-facing interfaces. This action is irreversible and should be used only when a doctor '
    'account is confirmed fraudulent or inactive.')
append_spacer(doc)

append_h3(doc, 'B.2.1.1  When to Remove a Doctor Account')
append_body(doc,
    'Removal is appropriate when: a doctor account has been flagged as a duplicate registration, '
    'when a practitioner has formally withdrawn from the platform, or when account credentials '
    'are suspected to be compromised and no remediation is possible. For temporary suspension, '
    'the recommendation is to coordinate a password reset rather than full deletion, preserving '
    'historical appointment and review data.')
append_spacer(doc)

# B.3
append_h1(doc, 'B.3  Managing Patients')
append_body(doc,
    'The Patients section lists all registered patient accounts retrieved via GET /admin/patients. '
    'The table displays each patient\'s full name, email address, gender, date of birth, and '
    'account creation date. Administrators may delete any patient account using the per-row '
    '"Delete" button, which calls DELETE /admin/patient/:id. The same confirmation dialog and '
    'data-retention considerations described for doctor removal apply here. Patient chat histories, '
    'appointment records, and reviews are retained for audit and historical analysis.')
append_spacer(doc)

append_h2(doc, 'B.3.1  Privacy Considerations')
append_body(doc,
    'Patient data is sensitive personal health information. Administrators must comply with '
    'applicable data protection regulations when viewing, exporting, or deleting patient records. '
    'The Admin Dashboard does not provide a bulk-export or bulk-delete function; all operations '
    'are performed on individual records with explicit confirmation steps to prevent accidental '
    'mass data loss. Access to the admin interface must be strictly limited to authorised '
    'personnel and should be performed over a secure, encrypted network connection.')
append_spacer(doc)

# B.4
append_h1(doc, 'B.4  Managing Appointments')
append_body(doc,
    'The Appointments section retrieves all appointment records via GET /admin/appointments and '
    'presents them in a filterable table showing appointment ID, patient name, doctor name, '
    'scheduled date and time, and current status. Administrators can delete individual appointments '
    'using DELETE /admin/appointment/:id. This capability is intended for removing test data, '
    'resolving data conflicts, or purging appointments created erroneously. Deleting an appointment '
    'does not automatically cancel the associated VideoSDK meeting room; the room expires '
    'naturally on VideoSDK\'s servers.')
append_spacer(doc)

append_h2(doc, 'B.4.1  Managing Schedules')
append_body(doc,
    'The Schedules section lists all doctor availability slots fetched via GET /admin/schedules. '
    'Each record shows the owning doctor, date, start time, end time, and status. Administrators '
    'may delete orphaned or erroneous schedule records using DELETE /admin/schedule/:id. Booked '
    'slots should only be deleted after confirming that the associated appointment has been '
    'handled appropriately to avoid leaving appointments in an inconsistent state.')
append_spacer(doc)

# B.5
append_h1(doc, 'B.5  Managing Meeting Notes and Reviews')
append_body(doc,
    'Meeting Notes records are accessible via GET /admin/meeting-notes and display the associated '
    'appointment, generation status (processing, complete, or failed), and a download link for '
    'completed PDF notes. Failed notes may indicate that audio was not captured during the '
    'consultation or that the transcription service encountered an error. DELETE '
    '/admin/meeting-note/:id removes the database record; the physical PDF file on the server '
    'should be removed separately if storage reclamation is required.')
append_spacer(doc)

append_h2(doc, 'B.5.1  Moderating Reviews')
append_body(doc,
    'The Reviews section lists all patient-submitted reviews via GET /admin/reviews, displaying '
    'the reviewer\'s name, the reviewed doctor, the star rating, the review text, and the '
    'submission date. Administrators may remove inappropriate, abusive, or fraudulent reviews '
    'using DELETE /admin/review/:id. Deletion of a review updates the associated doctor\'s '
    'aggregate rating dynamically through the rating recalculation logic in the backend.')
append_spacer(doc)

append_h3(doc, 'B.5.1.1  Criteria for Review Removal')
append_body(doc,
    'A review should be considered for removal if it: contains offensive or discriminatory language, '
    'references personal information that should not be publicly visible, appears to be a duplicate '
    'submission from the same appointment, or is confirmed to be submitted by a user who did not '
    'complete a genuine appointment with the rated doctor. Legitimate negative reviews must not be '
    'removed solely due to low ratings.')
append_spacer(doc)

page_break(doc)

# ══════════════════════════════════════════════════════════════════════════════
# APPENDIX C — INFORMATION / PROMOTIONAL MATERIAL
# ══════════════════════════════════════════════════════════════════════════════

append_title(doc, 'Appendix C: Information / Promotional Material')
append_spacer(doc)
append_body(doc,
    'This appendix presents the suite of promotional and informational materials developed to '
    'communicate the MediCare AI platform\'s value proposition to prospective patients, healthcare '
    'practitioners, and institutional stakeholders. The materials are designed to convey the '
    'platform\'s core benefits — accessible teleconsultation, AI-assisted health guidance, digital '
    'prescriptions, and automated meeting documentation — in concise, visually engaging formats '
    'suitable for both print distribution and digital display. Each asset targets a specific '
    'communication context, from a detailed informational brochure through to large-format outdoor '
    'banners, ensuring consistent brand messaging across all channels. The colour palette, '
    'typography, and tone used across all materials align with the MediCare AI brand identity: '
    'navy (#0A2540), teal (#00B4A0), and clean white backgrounds conveying clinical trust and '
    'technological modernity.')
append_spacer(doc)

# C.1 Brochure
append_h1(doc, 'C.1  Brochure')
append_body(doc,
    'The MediCare AI brochure is a tri-fold A4 print piece designed for distribution at hospitals, '
    'clinics, pharmacies, and university health centres. The front panel features the MediCare AI '
    'logo, the tagline "Your Health. Anytime. Anywhere.", and a high-quality illustration of a '
    'patient consulting a doctor via video call. The inner panels are structured across three '
    'columns: the first column introduces the platform and its mission of democratising access to '
    'qualified healthcare through technology; the second column details the five core features — '
    'one-click video consultations, AI symptom checking, digital prescriptions, automated meeting '
    'notes, and doctor discovery with verified reviews; the third column contains a step-by-step '
    '"Get Started" guide (scan QR code → register in 60 seconds → book your first appointment → '
    'consult from anywhere). The back panel provides the platform URL, support email, QR code '
    'linking to the registration page, and a short privacy assurance statement.')
append_spacer(doc)

append_h2(doc, 'C.1.1  Brochure Design Specifications')
append_body(doc,
    'Dimensions: 297 mm × 210 mm (A4, landscape), folded to 99 mm × 210 mm panels. Bleed: 3 mm '
    'on all sides. Resolution: 300 DPI minimum for print. Primary font: Calibri for headings, '
    'Times New Roman for body text. Primary colours: Navy (#0A2540) for headings and icon '
    'backgrounds, Teal (#00B4A0) for call-to-action buttons and accent lines, White (#FFFFFF) for '
    'panel backgrounds. File format for print submission: PDF/X-1a with embedded fonts.')
append_spacer(doc)

append_h3(doc, 'C.1.1.1  Key Messages')
append_body(doc,
    'The brochure communicates three primary messages: (1) Convenience — "See a verified doctor '
    'from your home in minutes, no travel required."; (2) Intelligence — "Our AI symptom advisor '
    'gives you evidence-based health guidance before your appointment."; (3) Continuity — '
    '"Digital prescriptions and AI-generated meeting notes keep your complete medical record in '
    'one secure place, accessible anytime."')
append_spacer(doc)

# C.2 Flyer
append_h1(doc, 'C.2  Flyer')
append_body(doc,
    'The MediCare AI flyer is a single-sided A5 promotional sheet intended for mass distribution '
    'at campus events, community health fairs, and waiting rooms. The design prioritises immediate '
    'visual impact: a bold teal header band containing the platform name and tagline, a central '
    'section listing the three most compelling user benefits as icon-accompanied bullet points '
    '(Video Consultation icon — "Consult verified doctors from your phone or laptop"; Brain/AI icon '
    '— "AI symptom checker available 24/7 at no extra cost"; Document icon — "Digital prescriptions '
    'delivered instantly after your appointment"), and a footer containing a large QR code and the '
    'URL. The overall layout is designed to be readable within three seconds, with a call to action '
    '"Scan to Book Your First Consultation" prominently placed below the QR code.')
append_spacer(doc)

append_h2(doc, 'C.2.1  Flyer Design Specifications')
append_body(doc,
    'Dimensions: 148 mm × 210 mm (A5, portrait). Resolution: 300 DPI minimum. The flyer is '
    'designed for single-sided digital printing on 130 gsm gloss paper. Colour mode: CMYK. '
    'The QR code encodes the direct patient registration URL and is tested for scannability at '
    'sizes down to 25 mm × 25 mm. A digital version optimised for social media (1080 × 1920 px '
    'at 72 DPI) is produced from the same artwork by removing bleed and adjusting margins.')
append_spacer(doc)

# C.3 Standee
append_h1(doc, 'C.3  Standee')
append_body(doc,
    'The MediCare AI pull-up standee is a 85 cm × 200 cm roll-up display banner designed for '
    'placement at health exhibitions, conference registration desks, and hospital reception areas. '
    'The upper third of the standee features the platform logo against a navy gradient background '
    'with the tagline in large white Calibri Bold text. The middle third presents four service '
    'pillars arranged horizontally with illustrative icons and two-line descriptions: "Doctor '
    'Discovery" (Search verified specialists by name or specialty), "Video Consultation" (HD '
    'video calls from any device, no app install needed), "AI Health Advisor" (Evidence-based '
    'symptom assessment powered by medical literature), and "Digital Prescription" (Structured '
    'prescriptions delivered to your patient account instantly). The lower third contains a QR '
    'code, the platform URL, and a short social proof statement citing the number of available '
    'specialists.')
append_spacer(doc)

append_h2(doc, 'C.3.1  Standee Design Specifications')
append_body(doc,
    'Dimensions: 850 mm × 2000 mm. Resolution: 150 DPI at full size (minimum). File format for '
    'production: PDF with 5 mm bleed. Material recommendation: 440 gsm PVC vinyl with matte '
    'laminate finish to reduce glare under indoor lighting. The standee includes a graphic safe '
    'zone of 50 mm from the bottom edge to account for the pull-up mechanism housing. Text '
    'elements are sized for legibility from a distance of 1.5 to 3 metres.')
append_spacer(doc)

append_h3(doc, 'C.3.1.1  Placement Guidelines')
append_body(doc,
    'For optimal visibility, the standee should be placed at eye level in high-foot-traffic areas '
    'such as entrances, waiting rooms, or near information desks. It should not be positioned '
    'directly in front of windows or strong light sources, as back-lighting significantly reduces '
    'text readability on matte vinyl. Two standees placed symmetrically at a booth entrance '
    'create a professional framing effect for exhibition contexts.')
append_spacer(doc)

# C.4 Banner
append_h1(doc, 'C.4  Banner')
append_body(doc,
    'The MediCare AI horizontal banner is designed for large-format outdoor and indoor display, '
    'such as stage backdrops at health conferences, building facades during launch events, or '
    'stretched across auditorium screens. At 300 cm × 90 cm, the banner communicates a single '
    'high-impact message: the MediCare AI logotype in large format on the left, the central '
    'message "Bringing Healthcare Closer — One Click at a Time" in white 72 pt Calibri Bold on '
    'a navy background, and on the right a QR code panel with "Book Your Consultation Now" in '
    'teal. The sparse, three-zone layout ensures the banner remains legible from distances of '
    'five metres or more and photographs cleanly for event documentation.')
append_spacer(doc)

append_h2(doc, 'C.4.1  Banner Design Specifications')
append_body(doc,
    'Dimensions: 3000 mm × 900 mm (landscape). Resolution: 100 DPI at full size. Material: '
    '510 gsm PVC flex with reinforced hemming and brass eyelets at 500 mm intervals for '
    'suspension. Colour mode: CMYK with Pantone equivalents — Navy: Pantone 2965 C; Teal: '
    'Pantone 3272 C. File format for production: PDF/X-4 with all fonts embedded and images '
    'flattened. A secondary indoor version (200 cm × 60 cm, 150 DPI, fabric substrate) is '
    'produced for conference exhibitor booths.')
append_spacer(doc)

page_break(doc)

# ══════════════════════════════════════════════════════════════════════════════
# APPENDIX D — INSTALLATION AND DEPLOYMENT GUIDE
# ══════════════════════════════════════════════════════════════════════════════

append_title(doc, 'Appendix D: Installation and Deployment Guide')
append_spacer(doc)
append_body(doc,
    'This appendix provides a technical guide for setting up and deploying the MediCare AI '
    'telemedicine platform in a local development environment and a production server environment. '
    'It covers the prerequisites for each of the three service tiers — Frontend, Backend, and AI '
    'Service — including dependency installation, environment variable configuration, and startup '
    'procedures. The guide also describes the build and deployment steps required to serve the '
    'platform in a production context with appropriate security hardening. Readers are expected '
    'to have basic familiarity with command-line interfaces, Node.js package management, and '
    'Python virtual environments. All commands are presented for Unix-like environments; Windows '
    'equivalents are noted where they differ.')
append_spacer(doc)

# D.1
append_h1(doc, 'D.1  Prerequisites and System Requirements')
append_body(doc,
    'Before installation, ensure the target machine meets the following minimum requirements: '
    '8 GB RAM (16 GB recommended for running the AI transcription model), a modern multi-core '
    'CPU (Intel i5 8th generation or equivalent), and at least 10 GB of free disk space for '
    'dependencies, FAISS vector index, and model weights. The following software must be '
    'pre-installed: Node.js v20 LTS or later (with npm v10+), Python 3.11 or later, uv package '
    'manager (pip install uv), MongoDB 7.0 or later (local instance or MongoDB Atlas cloud URI), '
    'and Git for cloning the repository.')
append_spacer(doc)

append_h2(doc, 'D.1.1  Required API Keys and Environment Variables')
append_body(doc,
    'The following external service credentials must be obtained before deployment: '
    'GROQ_API_KEY — obtained from console.groq.com; HF_TOKEN — obtained from '
    'huggingface.co (required for downloading embedding model weights); VIDEOSDK_API_KEY and '
    'VIDEOSDK_SECRET — obtained from app.videosdk.live; JWT_SECRET and P_JWT_SECRET — '
    'user-defined random strings of at least 32 characters for signing doctor and patient '
    'JWT tokens respectively; MONGODB_URI — the MongoDB connection string. All backend '
    'variables are placed in a .env file in the BackEnd/ directory. AI service variables '
    'are placed in a .env file in the AI/ directory.')
append_spacer(doc)

append_h3(doc, 'D.1.1.1  Sample Environment File Structure')
append_body(doc,
    'BackEnd/.env should contain: MONGODB_URI, JWT_SECRET (maps to SecretKey), P_JWT_SECRET '
    '(maps to P_SecretKey), VIDEOSDK_API_KEY, and VIDEOSDK_SECRET. AI/.env should contain: '
    'GROQ_API_KEY and HF_TOKEN. FrontEnd/.env should contain: VITE_API_URL pointing to the '
    'backend base URL (e.g., http://localhost:4000 for development or the production HTTPS '
    'URL). Never commit .env files to version control; ensure they are listed in .gitignore.')
append_spacer(doc)

# D.2
append_h1(doc, 'D.2  Backend Installation and Startup')
append_body(doc,
    'Clone the repository and navigate to the BackEnd/ directory. Run "npm install" to install '
    'all Express 5 and Mongoose dependencies. Ensure the MongoDB service is running and the '
    'MONGODB_URI in .env is correctly configured. Start the development server with "node '
    'index.js"; the server listens on port 4000 by default. For production, it is recommended '
    'to use a process manager such as PM2 ("pm2 start index.js --name medicare-backend") and '
    'place the backend behind an Nginx reverse proxy with TLS termination to serve HTTPS traffic '
    'on port 443.')
append_spacer(doc)

append_h2(doc, 'D.2.1  Frontend Installation and Build')
append_body(doc,
    'Navigate to the FrontEnd/ directory and run "npm install" to install all React 19, Vite, '
    'and TailwindCSS v4 dependencies. For development, run "npm run dev" to start the Vite '
    'development server on port 5173 with hot module replacement enabled. For production, run '
    '"npm run build" to generate an optimised static bundle in the dist/ folder. Deploy the '
    'contents of dist/ to a static file host (Nginx, Vercel, Netlify, or AWS S3 + CloudFront). '
    'Ensure VITE_API_URL in FrontEnd/.env points to the production backend HTTPS URL before '
    'running the production build.')
append_spacer(doc)

append_h2(doc, 'D.2.2  AI Service Installation and Vector Store Initialisation')
append_body(doc,
    'Navigate to the AI/ directory. Run "uv sync" to create a virtual environment and install '
    'all Python dependencies declared in pyproject.toml, including FastAPI, LangChain, '
    'langchain-groq, faster-whisper, and faiss-cpu. Place the five source medical PDF documents '
    'in the AI/data/ directory. Run "python create_memory_for_llm.py" to parse the PDFs, '
    'generate HuggingFace sentence-transformer embeddings, and persist the FAISS vector index '
    'at AI/vectorstore/db_faiss/. This step is required only once, or whenever the source '
    'documents change. Start the FastAPI server with "uv run uvicorn API.main:app --reload" '
    'to listen on port 8000.')
append_spacer(doc)

append_h3(doc, 'D.2.2.1  Verifying the AI Service')
append_body(doc,
    'After starting the FastAPI server, navigate to http://localhost:8000/docs in a browser to '
    'access the auto-generated Swagger UI. Use the /chat endpoint\'s "Try it out" feature to '
    'send a test message with an arbitrary session_id and a medical question. A successful '
    'response includes an "answer" field with LLM-generated text and a "sources" field listing '
    'the retrieved document passages. If the response returns a 503 error, confirm that the '
    'MedicalRAG engine initialised without errors in the console output; failures typically '
    'indicate a missing GROQ_API_KEY or an absent FAISS vector store.')
append_spacer(doc)

# ── Save ─────────────────────────────────────────────────────────────────────
output_path = 'MediCare_AI_Appendices.docx'
doc.save(output_path)
print(f"Done! Saved: {output_path}")
print(f"Paragraphs written: {len(doc.paragraphs)}")
